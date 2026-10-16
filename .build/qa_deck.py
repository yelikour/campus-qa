# -*- coding: utf-8 -*-
"""代码 QA:文本溢出估算 / 越界 / 文本框重叠检查(近似启发式)。"""
import sys
from pptx import Presentation
from pptx.util import Emu

PATH = sys.argv[1] if len(sys.argv) > 1 else "ppt/智问校园CampusMind项目讲解PPT.pptx"
prs = Presentation(PATH)
SW, SH = prs.slide_width, prs.slide_height
EMU_IN = 914400

def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
        else:
            yield sh

issues = 0
for si, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in walk(slide.shapes):
        if not getattr(sh, "has_text_frame", False) or not sh.text_frame.text.strip():
            continue
        w_in, h_in = sh.width / EMU_IN, sh.height / EMU_IN
        x_in, y_in = sh.left / EMU_IN, sh.top / EMU_IN
        # 越界
        if x_in < -0.01 or y_in < -0.01 or x_in + w_in > SW / EMU_IN + 0.01 or y_in + h_in > SH / EMU_IN + 0.01:
            print(f"[越界] slide{si} '{sh.text_frame.text[:14]}' box=({x_in:.2f},{y_in:.2f},{w_in:.2f},{h_in:.2f})")
            issues += 1
        # 溢出估算:CJK 全角按 size*1.05pt 宽,拉丁按 0.62
        total_h = 0.0
        for p in sh.text_frame.paragraphs:
            text = "".join(r.text for r in p.runs)
            if not text:
                total_h += 0.12
                continue
            size = max((r.font.size.pt if r.font.size else 14) for r in p.runs)
            cjk = sum(1 for c in text if ord(c) > 0x2E80)
            lat = len(text) - cjk
            line_w_pt = size * 1.05 * cjk + size * 0.62 * lat
            avail_pt = max(w_in * 72 - 6, 10)
            lines = max(1, -(-line_w_pt // avail_pt))
            ls = 1.2
            total_h += lines * size * ls / 72
        if total_h > h_in * 1.28 + 0.06:
            print(f"[溢出?] slide{si} '{sh.text_frame.text[:16]}' est={total_h:.2f}in box_h={h_in:.2f}in")
            issues += 1
        boxes.append((x_in, y_in, w_in, min(h_in, total_h), sh.text_frame.text[:12]))
    # 重叠(同层文本框,忽略 30% 以下重叠)
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            a, b = boxes[i], boxes[j]
            ox = max(0, min(a[0]+a[2], b[0]+b[2]) - max(a[0], b[0]))
            oy = max(0, min(a[1]+a[3], b[1]+b[3]) - max(a[1], b[1]))
            if ox * oy > 0.05 and ox * oy / min(a[2]*a[3], b[2]*b[3]) > 0.3:
                print(f"[重叠] slide{si} '{a[4]}' x '{b[4]}' area={ox*oy:.2f}")
                issues += 1
print("QA issues:", issues)
